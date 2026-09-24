"""Literal workbook validation inputs; both ambiguous layouts preserved."""
import json,csv,hashlib,shutil,datetime
from pathlib import Path
import openpyxl
from rdkit import Chem
from rdkit.Chem import Descriptors
R=Path(__file__).resolve().parents[1];D=Path('/mnt/r/plastchem-euler/phase8-v1')
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
book=Path('/home/aaltamimi2/langchain-STRAP-v10-core/data/zhou_contamintant_removal_SI_Data.xlsx')
assert sha(book)=='e5640e458d7949d0a3570143518b18aa14b8cb17b1453fcf78606d4aef26fcde'
wb=openpyxl.load_workbook(book,data_only=True);names=['BBP','DBP','DEHP','DEP','DiDP','DiNP','DnHP','DnOP']
keys=['IRIAEXORFWYRCZ','DOIRQSBPFJWKBE','BJQHLKABXJIVAM','FLKPEMZONWLCSK','ZVFDTKUVRCTHQE','HBGGXOJOCNVPFY','KCXZNSGUUQJJTR','MQIUGAXCHLFZKX']
solutes={}
for name,key in zip(names,keys):
 record=next((R/'state/campaign-v1/records').glob(key+'*.json'));r=json.loads(record.read_text());assert r['status']=='converged';surface=Path(r['archive_path'])/'surface.orcacosmo';h=sha(surface);assert h==r['surface_sha256'];rel='surfaces/'+h+'.orcacosmo';shutil.copyfile(surface,D/rel)
 solutes[name]={'B':rel,'input':r['input'],'record_sha256':sha(record),'surface_sha256':h,'cpu_model':r['cpu_model'],'molecular_weight_g_mol':float(r['input']['molecular_weight_g_mol'])}
mapping={r['solvent_raw'].strip().lower():r['solvent_key'] for r in csv.DictReader((R/'reports/phase8-0-inventory/product-solvent-key-map.csv').open())}
registry=json.loads((R/'state/thermodynamics-v1/library-registry.json').read_text());masses={}
# Solvent masses are re-derived from the pinned identity SMILES using RDKit.
solvent_smiles={}
for s in registry['solvents']:
 if s['status']!='ready':continue
 rec=json.loads(Path(s['source_record']).read_text());i=rec.get('input',{});smiles=i.get('smiles') or rec.get('smiles');mol=Chem.MolFromSmiles(smiles);assert mol is not None
 solvent_smiles[s['solvent_key']]={'smiles':smiles,'record':s['source_record'],'identity':s['source_identity'],'molecular_weight_g_mol':Descriptors.MolWt(mol),'mass_method':'RDKit average atomic weights, neutral pinned solvent SMILES'}
logs=[];misc=[];units=[];ws=wb['Phthalates_log_D']
for row in range(2,34):
 solvent=mapping[ws.cell(row,1).value.strip().lower()]
 for j,name in enumerate(names):logs.append({'solute':name,'solvent':solvent,'cell':ws.cell(row,j+3).coordinate,'value':float(ws.cell(row,j+3).value),'temperature_K':298.15,'polymer':'pvc'})
ws=wb['Phthalates_Miscibility']
for row in range(3,35):
 solvent=mapping[ws.cell(row,1).value.strip().lower()];high=float(ws.cell(row,3).value)+273.15
 for j,name in enumerate(names):
  for regime,T in [('RT',298.15),('high',high)]:
   units.append({'solute':name,'solvent':solvent,'regime':regime,'temperature_K':T,'literal_high_temperature_C':high-273.15})
   for layout,col in [('blocked',4+j+(8 if regime=='high' else 0)),('paired',4+2*j+(1 if regime=='high' else 0))]:
    cell=ws.cell(row,col);misc.append({'solute':name,'solvent':solvent,'regime':regime,'temperature_K':T,'layout':layout,'cell':cell.coordinate,'literal_header':ws.cell(1,col).value,'literal_regime_header':ws.cell(2,col).value,'literal_verdict':cell.value})
out={'utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),'workbook':str(book),'workbook_sha256':sha(book),'solutes':solutes,'solvent_identities':solvent_smiles,'logP_reference':logs,'miscibility_reference':misc,'lle_units':units,'literal_conflicts':list(csv.DictReader((R/'reports/phase8-0-inventory/phthalate-duplicate-regime-keys.csv').open()))}
(D/'validation-inputs.json').write_text(json.dumps(out,indent=2)+'\n');print('Prepared',len(logs),'logP and',len(units),'LLE systems;',len(misc),'layout reference rows')
