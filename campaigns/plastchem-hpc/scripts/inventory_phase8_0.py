"""A-8 Phase 8.0: file/table inventory only; no thermodynamic engine or scheduler calls."""
from pathlib import Path
import collections, csv, datetime, hashlib, json, subprocess
import duckdb
import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'reports/phase8-0-inventory'
BULK = Path('/mnt/r/plastchem-euler/phase8-0-inventory-2026-09-23')
PRODUCT = Path('/home/aaltamimi2/dissolve-v12-builder-1')
WORKBOOK = Path('/home/aaltamimi2/langchain-STRAP-v10-core/data/zhou_contamintant_removal_SI_Data.xlsx')
PAPER = Path('/home/aaltamimi2/dissolve-v12-work/contaminant.pdf')
PIN = 'e5640e458d7949d0a3570143518b18aa14b8cb17b1453fcf78606d4aef26fcde'

def sha(path):
    h = hashlib.sha256()
    with path.open('rb') as f:
        for block in iter(lambda: f.read(1048576), b''):
            h.update(block)
    return h.hexdigest()

def save(path, obj):
    path.write_text(json.dumps(obj, indent=2, default=str) + '\n')

def csvout(path, rows):
    if not rows:
        return
    with path.open('w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]), lineterminator='\n')
        w.writeheader()
        w.writerows(rows)

def query(con, sql):
    c = con.execute(sql)
    names = [x[0] for x in c.description]
    return [dict(zip(names, row)) for row in c.fetchall()]

def main():
    OUT.mkdir(parents=True, exist_ok=True)
    BULK.mkdir(parents=True, exist_ok=True)
    started = datetime.datetime.now(datetime.timezone.utc).isoformat()
    assert sha(WORKBOOK) == PIN
    sources = [WORKBOOK, PAPER, ROOT/'CHARTER.txt',
               PRODUCT/'src/dissolve/contaminants.py', PRODUCT/'src/dissolve/thermodynamics.py',
               PRODUCT/'src/dissolve/data/contaminants.duckdb',
               PRODUCT/'src/dissolve/data/thermodynamics.duckdb',
               ROOT/'scripts/thermodynamic_prediction.py', ROOT/'scripts/a7_phased_unit.py',
               Path('/home/aaltamimi2/langchain-STRAP-v10-core/src/strap/services/contaminant_data_service.py')]
    package = Path('/home/aaltamimi2/.venvs/cosmo-logp/lib/python3.11/site-packages/opencosmorspy')
    sources.extend(package / name for name in ['cosmors.py', 'parameterization.py', 'molecules.py'])
    pins = [{'path': str(p), 'size': p.stat().st_size, 'sha256': sha(p)} for p in sources]
    wb = openpyxl.load_workbook(WORKBOOK, read_only=False, data_only=False)
    cached = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    sheets, columns, cells, values, labels = [], [], [], [], []
    for ws in wb:
        pfas = ws.title.startswith('PFAS')
        misc = 'Miscibility' in ws.title
        start = 3 if pfas or misc else 2
        name_row = 2 if pfas else 1
        solvent_col = 2 if pfas else 1
        first_value = 3 if pfas else 4 if misc else 3
        data_rows = [r for r in range(start, ws.max_row+1)
                     if isinstance(ws.cell(r, first_value).value, (int, float))
                     or ws.cell(r, first_value).value in ['Yes', 'No']]
        count = 0
        for row in ws:
            for cell in row:
                if cell.value is not None:
                    cells.append({'sheet': ws.title, 'cell': cell.coordinate,
                                  'value': cell.value, 'cached_value': cached[ws.title][cell.coordinate].value,
                                  'type': cell.data_type})
        for col in range(1, ws.max_column+1):
            vs = [ws.cell(r, col).value for r in range(1, ws.max_row+1)]
            role = 'empty_formatted_column' if all(x is None for x in vs) else (
                'contaminant_value' if col >= first_value and ws.cell(name_row,col).value else 'identifier_or_metadata')
            columns.append({'sheet':ws.title,'column':get_column_letter(col),'column_number':col,
                'row1_literal':ws.cell(1,col).value,'row2_literal':ws.cell(2,col).value,
                'role':role,'nonempty_cells':sum(x is not None for x in vs),
                'data_count':sum(ws.cell(r,col).value is not None for r in data_rows),
                'data_value_types':','.join(sorted({type(ws.cell(r,col).value).__name__ for r in data_rows})),
                'literal_regime':ws.cell(2,col).value if misc and not pfas and col>=4 else None})
        for r in data_rows:
            solvent = str(ws.cell(r,solvent_col).value).strip()
            labels.append({'sheet':ws.title,'excel_row':r,'solvent_raw':solvent,
                           'boiling_point_C':ws.cell(r,2).value if not pfas else None,
                           'T_higher_C':ws.cell(r,3).value if misc and not pfas else None})
            for col in range(first_value,ws.max_column+1):
                name = ws.cell(name_row,col).value
                value = ws.cell(r,col).value
                if name is None or value is None:continue
                values.append({'sheet':ws.title,'cell':ws.cell(r,col).coordinate,
                    'contaminant_literal':str(name).strip(),'solvent_raw':solvent,
                    'regime_literal':ws.cell(2,col).value if misc and not pfas else None,
                    'value':value})
                count += 1
        sheets.append({'sheet':ws.title,'declared_rows':ws.max_row,'declared_columns':ws.max_column,
                       'solvent_rows':len(data_rows),'value_cells':count,
                       'merged_ranges':[str(x) for x in ws.merged_cells.ranges],
                       'formulas':sum(c.data_type=='f' for row in ws for c in row),
                       'nonempty_cells':sum(c.value is not None for row in ws for c in row)})
    csvout(OUT/'workbook-columns.csv',columns)
    csvout(OUT/'workbook-solvent-temperature-rows.csv',labels)
    csvout(BULK/'workbook-all-nonempty-cells.csv',cells)
    csvout(BULK/'workbook-data-values-literal.csv',values)
    save(OUT/'workbook-sheets.json',sheets)
    c = duckdb.connect(str(PRODUCT/'src/dissolve/data/contaminants.duckdb'),read_only=True)
    c.execute("SET threads=1")
    c.execute("SET memory_limit='256MB'")
    tables = {}
    for (t,) in c.execute('SHOW TABLES').fetchall():
        rows=query(c,f'SELECT * FROM "{t}"')
        tables[t]={'columns':query(c,f'DESCRIBE "{t}"'),'count':len(rows)}
        csvout(BULK/f'product-{t}.csv',rows)
    q = {
      'logd_by_family': 'SELECT family,count(*) AS n,count(distinct contaminant_key) AS contaminants,count(distinct solvent_key) AS solvents,count(logd) AS nonnull_values FROM logd GROUP BY family',
      'miscibility_regimes':'SELECT family,temperature_regime,count(*) AS n,min(temperature_c) AS min_C,max(temperature_c) AS max_C FROM miscibility GROUP BY ALL',
      'phthalate_regime_counts':"SELECT contaminant,temperature_regime,count(*) AS n FROM miscibility WHERE family='Phthalates' GROUP BY ALL ORDER BY contaminant,temperature_regime",
      'duplicate_miscibility_keys':'SELECT family,contaminant_key,solvent_key,temperature_regime,count(*) AS n,count(distinct miscible) AS distinct_verdicts FROM miscibility GROUP BY ALL HAVING count(*)>1',
      'solvent_names':'SELECT DISTINCT solvent_raw,solvent_key,solvent_normalized FROM logd ORDER BY solvent_key',
      'metadata':'SELECT * FROM metadata',
      'identity_aliases':'SELECT DISTINCT contaminant_key,canonical_name,family,cas_number,resolution_basis,pubchem_cid FROM contaminant_aliases ORDER BY contaminant_key',
    }
    observations = {k:query(c,sql) for k,sql in q.items()}
    def norm(value):
        return ' '.join(str(value).strip().casefold().split())
    reconciliation = {}
    for kind in ['logd', 'miscibility']:
        expected, actual = [], []
        for row in values:
            if (kind == 'miscibility') != ('Miscibility' in row['sheet']):
                continue
            regime = norm(row['regime_literal']).replace(' ', '_') if row['regime_literal'] else 'unspecified'
            value = norm(row['value']) == 'yes' if kind == 'miscibility' else float(row['value'])
            expected.append((norm(row['contaminant_literal']), norm(row['solvent_raw']),
                             regime if kind == 'miscibility' else None, value))
        for row in query(c, 'SELECT * FROM ' + kind):
            actual.append((norm(row['contaminant']), norm(row['solvent_raw']),
                           row['temperature_regime'] if kind == 'miscibility' else None,
                           row['miscible'] if kind == 'miscibility' else row['logd']))
        a, b = collections.Counter(expected), collections.Counter(actual)
        reconciliation[kind] = {'source_cells': len(expected), 'database_rows': len(actual),
                                'unmatched_source': sum((a-b).values()), 'unmatched_database': sum((b-a).values())}
        assert a == b, 'Literal workbook/database discrepancy: ' + kind
    observations['literal_workbook_database_reconciliation'] = reconciliation
    csvout(OUT/'phthalate-duplicate-regime-keys.csv',observations['duplicate_miscibility_keys'])
    csvout(OUT/'product-solvent-key-map.csv',observations['solvent_names'])
    csvout(OUT/'product-contaminant-identities.csv',observations['identity_aliases'])
    c.close()
    c=duckdb.connect(str(PRODUCT/'src/dissolve/data/thermodynamics.duckdb'),read_only=True)
    c.execute('SET threads=1')
    observations['grid_schema']=query(c,'DESCRIBE solubility_grid')
    observations['grid_counts']=query(c,'SELECT count(*) AS rows,count(distinct (polymer,solvent)) AS pairs,count(distinct polymer) AS polymers,count(distinct solvent) AS solvents,min(temperature_c) AS min_C,max(temperature_c) AS max_C FROM solubility_grid')
    observations['grid_filter_counts']=query(c,'SELECT is_valid,invalid_reason,count(*) AS n FROM solubility_grid GROUP BY ALL')
    observations['grid_source_counts']=query(c,'SELECT source_table,count(*) AS n FROM solubility_grid GROUP BY ALL')
    observations['grid_temperatures']=query(c,'SELECT DISTINCT temperature_c FROM solubility_grid ORDER BY temperature_c')
    observations['grid_sources']=query(c,'SELECT * FROM asset_sources')
    c.close()
    summary={'started_utc':started,'finished_utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),
             'phase':'8.0 inventory only','no_thermodynamic_calculations':True,
             'sources':pins,'workbook_sheets':sheets,'database_tables':tables,'observations':observations,
             'bulk_evidence_directory':str(BULK)}
    for pin in pins:
        assert sha(Path(pin['path']))==pin['sha256'], 'Source changed during inventory: '+pin['path']
    summary['all_source_digests_unchanged_after_read']=True
    save(OUT/'inventory-evidence.json',summary)
    save(OUT/'source-pins.json',pins)
    save(OUT/'bulk-evidence-manifest.json', [{'path':str(p),'bytes':p.stat().st_size,'sha256':sha(p)}
                                          for p in sorted(BULK.iterdir()) if p.is_file()])
    print(json.dumps({'sheets':sheets,'duplicate_keys':len(observations['duplicate_miscibility_keys']),
                      'contradictory_duplicate_keys':sum(r['distinct_verdicts']>1 for r in observations['duplicate_miscibility_keys']),
                      'grid_counts':observations['grid_counts']}))

if __name__=='__main__':main()
